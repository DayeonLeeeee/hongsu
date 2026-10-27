# -*- coding: utf-8 -*-
"""
rv2_codebook.py — R v2 자료(오세준 교수님 엑셀)를 서버 시작 시 직접 읽는다.

원칙 (README 이식 규칙 1·2):
  - 코드북·문항·줄별 기준판정·판정규칙은 전부 엑셀에서 읽는다. 서버 코드에 하드코딩하지 않는다.
  - 엑셀이 바뀌면 서버 재시작(또는 reload())만 하면 된다. R12·R13 확장도 엑셀만 바꾸면 됨.
  - 문항별 허용 코드 = 01시트 주_R코드 + 06시트 wrong일_때에 등장하는 코드 (키트 allowed_codes()와 동일).

엑셀 경로: 환경변수 RV2_XLSX (기본 materials/dataset_rv2.xlsx)
"""
import hashlib
import os
import re

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
XLSX_PATH = os.environ.get("RV2_XLSX", os.path.join(BASE_DIR, "materials", "dataset_rv2.xlsx"))

SHEETS = {
    "problems": "01_PROBLEMS",
    "r_codebook": "02_R_CODEBOOK",
    "j_codebook": "02b_J_CODEBOOK",
    "input_status": "02c_INPUT_STATUS",
    "line_judgment": "06_LINE_JUDGMENT",
    "prompt_rules": "09_PROMPT_RULES",
}

_mats = None


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v)


def _read_sheet(wb, name: str) -> list:
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [_cell(h) for h in rows[0]]
    out = []
    for r in rows[1:]:
        if r is None or all(c is None for c in r):
            continue
        out.append({header[i]: _cell(r[i]) if i < len(r) else "" for i in range(len(header)) if header[i]})
    return out


def load(path: str = None) -> dict:
    """엑셀을 읽어 키트 materials.json과 같은 구조의 dict를 만든다."""
    import openpyxl
    path = path or XLSX_PATH
    with open(path, "rb") as f:
        digest = hashlib.sha256(f.read()).hexdigest()[:8]
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    missing = [s for s in SHEETS.values() if s not in wb.sheetnames]
    if missing:
        raise ValueError(f"엑셀에 필요한 시트가 없음: {missing}")
    mats = {key: _read_sheet(wb, sheet) for key, sheet in SHEETS.items()}
    wb.close()
    mats["codebook_version"] = "R-v2"
    mats["xlsx_sha8"] = digest          # 재현성: 어떤 엑셀로 채점했는지
    mats["xlsx_path"] = path
    return mats


def get() -> dict:
    global _mats
    if _mats is None:
        _mats = load()
    return _mats


def reload() -> dict:
    global _mats
    _mats = load()
    return _mats


def problem(problem_id: str) -> dict:
    for p in get()["problems"]:
        if p.get("problem_id") == problem_id:
            return p
    raise KeyError(f"problem_id not found: {problem_id}")


def has_problem(problem_id: str) -> bool:
    return any(p.get("problem_id") == problem_id for p in get()["problems"])


def lines_for(problem_id: str) -> list:
    return [r for r in get()["line_judgment"] if r.get("problem_id") == problem_id]


def allowed_codes(problem_id: str) -> list:
    """키트 prompt_builder.allowed_codes()와 동일."""
    codes = set(re.findall(r"R\d+", problem(problem_id).get("주_R코드", "")))
    for r in lines_for(problem_id):
        codes |= set(re.findall(r"R\d+", r.get("wrong일_때", "") or ""))
    return sorted(codes, key=lambda c: int(c[1:]))


def problems_for_server_db() -> list:
    """서버 _problems_db 형식으로 변환 (기존 Q01~Q10 하드코딩을 대체)."""
    out = []
    for p in get()["problems"]:
        out.append({
            "id": p["problem_id"],
            "unit": "역함수",
            "type": "서술형",
            "title": f"역함수 {p.get('번호', '')}번 (난도 {p.get('난도', '')})",
            "statement": p.get("문항", ""),
            "correct_answer": p.get("정답", ""),
            "required_reasoning": p.get("required_evidence", ""),
            "expected_errors": re.findall(r"R\d+", p.get("주_R코드", "")),
            "source": p.get("출처", ""),
            "codebook": "R-v2",
            "concept_tags": ["역함수"],
            "chapter": "함수와 그래프",
            "score": 100,
            "rubric": [], "model_answer": "", "feedback_example": "", "grading_rules": "",
        })
    return out
